"""
Export Service — Lead Agent

Generates CSV and XLSX exports for lead batches.
Supports single batch, selected batches, and all-leads exports.
"""

import csv
import uuid
from pathlib import Path
from datetime import datetime, timezone

import openpyxl

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from platform_app.models.export_job import ExportJob, ExportStatus
from platform_app.models.lead import Lead
from platform_app.models.lead_batch import LeadBatch
from platform_app.config.settings import get_settings


# Columns for export
EXPORT_COLUMNS = [
    ("batch_name", "Batch"),
    ("business_name", "Business Name"),
    ("contact_name", "Contact Name"),
    ("designation", "Designation"),
    ("emails", "Emails"),
    ("phones", "Phones"),
    ("website", "Website"),
    ("address", "Address"),
    ("city", "City"),
    ("state", "State"),
    ("country", "Country"),
    ("matched_keyword", "Keyword"),
    ("source_primary", "Source"),
    ("rating", "Rating"),
    ("review_count", "Reviews"),
    ("linkedin_url", "LinkedIn"),
    ("instagram_url", "Instagram"),
    ("facebook_url", "Facebook"),
    ("x_url", "X (Twitter)"),
]


class ExportService:
    def __init__(self, session: AsyncSession, company_id: str | None = None):
        self.session = session
        self.company_id = company_id
        self.settings = get_settings()

    async def export_batch(self, batch_id: str, fmt: str = "csv") -> ExportJob:
        """Export a single batch."""
        batch = await self._get_batch(batch_id)
        if not batch:
            raise ValueError(f"Batch {batch_id} not found")

        leads = await self._get_leads_by_batches([batch_id])
        return await self._create_export(
            leads=leads,
            batch_names={batch_id: batch.name},
            fmt=fmt,
            selection_type="single_batch",
            batch_ids=[batch_id],
        )

    async def export_selected(
        self, batch_ids: list[str], fmt: str = "csv"
    ) -> ExportJob:
        """Export selected batches."""
        batches = {}
        for bid in batch_ids:
            batch = await self._get_batch(bid)
            if batch:
                batches[bid] = batch.name

        if not batches:
            raise ValueError("No valid batches found")

        leads = await self._get_leads_by_batches(list(batches.keys()))
        return await self._create_export(
            leads=leads,
            batch_names=batches,
            fmt=fmt,
            selection_type="selected_batches",
            batch_ids=list(batches.keys()),
        )

    async def export_all(self, fmt: str = "csv", unique_only: bool = False) -> ExportJob:
        """Export all leads."""
        stmt = select(Lead).where(Lead.company_id == self.company_id).order_by(Lead.created_at.desc())
        result = await self.session.execute(stmt)
        leads = list(result.scalars().all())

        # Build batch name map
        batch_ids = list({l.batch_id for l in leads})
        batch_names = {}
        for bid in batch_ids:
            batch = await self._get_batch(bid)
            if batch:
                batch_names[bid] = batch.name

        if unique_only:
            leads = self._deduplicate_leads(leads)

        return await self._create_export(
            leads=leads,
            batch_names=batch_names,
            fmt=fmt,
            selection_type="all",
            batch_ids=batch_ids,
        )

    async def _create_export(
        self,
        leads: list[Lead],
        batch_names: dict[str, str],
        fmt: str,
        selection_type: str,
        batch_ids: list[str],
    ) -> ExportJob:
        """Generate the export file and create a job record."""
        # Prepare data
        headers = [col[1] for col in EXPORT_COLUMNS]
        rows = []
        for lead in leads:
            row = []
            for col_key, _ in EXPORT_COLUMNS:
                if col_key == "batch_name":
                    row.append(batch_names.get(lead.batch_id, "Unknown"))
                elif col_key in ("emails", "phones"):
                    val = getattr(lead, col_key, [])
                    if not val:
                        val = []
                    row.append(", ".join(val) if isinstance(val, list) else str(val))
                else:
                    val = getattr(lead, col_key, "")
                    row.append(str(val) if val is not None else "")
            rows.append(row)

        # Generate file
        export_dir = self.settings.storage_root / "exports"
        export_dir.mkdir(parents=True, exist_ok=True)

        slug = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_name = f"leads_{slug}_{uuid.uuid4().hex[:6]}.{fmt}"
        file_path = export_dir / file_name

        if fmt == "csv":
            self._generate_csv(file_path, headers, rows)
        elif fmt == "xlsx":
            self._generate_xlsx(file_path, headers, rows)
        else:
            raise ValueError(f"Unsupported format: {fmt}")

        # Create job record
        job = ExportJob(
            company_id=self.company_id,
            export_type=fmt,
            selection_type=selection_type,
            selected_batch_ids=batch_ids,
            lead_count=len(leads),
            file_path=str(file_path),
            file_name=file_name,
            status=ExportStatus.COMPLETED.value,
            completed_at=datetime.now(timezone.utc),
        )
        self.session.add(job)
        await self.session.commit()
        await self.session.refresh(job)

        return job

    async def _get_batch(self, batch_id: str) -> LeadBatch | None:
        stmt = select(LeadBatch).where(
            LeadBatch.id == batch_id,
            LeadBatch.company_id == self.company_id
        )
        result = await self.session.execute(stmt)
        return result.scalar_one_or_none()

    async def _get_leads_by_batches(self, batch_ids: list[str]) -> list[Lead]:
        stmt = (
            select(Lead)
            .where(
                Lead.batch_id.in_(batch_ids),
                Lead.company_id == self.company_id
            )
            .order_by(Lead.batch_id, Lead.created_at.desc())
        )
        result = await self.session.execute(stmt)
        return list(result.scalars().all())

    def _deduplicate_leads(self, leads: list[Lead]) -> list[Lead]:
        """Deduplicate leads across batches using dedupe_key."""
        seen = set()
        unique = []
        for lead in leads:
            key = lead.dedupe_key or lead.id
            if key not in seen:
                seen.add(key)
                unique.append(lead)
        return unique

    def _generate_csv(self, file_path: Path, headers: list[str], data: list[list[str]]) -> None:
        with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            writer.writerows(data)

    def _generate_xlsx(self, file_path: Path, headers: list[str], data: list[list[str]]) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet()
            
        ws.title = "Leads"
        
        # Header styling
        from openpyxl.styles import Font, PatternFill, Alignment
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")

        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row in data:
            ws.append(row)

        # Auto-fit column widths (approximate)
        from openpyxl.utils import get_column_letter
        for col_idx, header in enumerate(headers, 1):
            max_len = len(header)
            for row in data:
                if col_idx - 1 < len(row):
                    max_len = max(max_len, len(row[col_idx - 1]))
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = min(max_len + 2, 40) # type: ignore

        wb.save(file_path)
